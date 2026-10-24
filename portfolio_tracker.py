#!/usr/bin/env python3
"""
portfolio_tracker.py

Reads an Excel file of holdings (Ticker, Shares, CostBasis), pulls the most
recent closing price for each ticker from Yahoo Finance, and writes a report
with market value and gain/loss per position and for the whole portfolio.

Input columns (headers are case-insensitive):
    Ticker      e.g. AAPL
    Shares      number of shares held
    CostBasis   TOTAL dollars paid for the position (not per share).
                Set COST_BASIS_IS_PER_SHARE = True below if your sheet
                holds per-share cost.

How to run (no terminal needed):
    Install once:  pip install pandas openpyxl yfinance
    Then open this file in your editor (Spyder, VS Code, PyCharm, IDLE...),
    click Run, and pick your holdings .xlsx in the window that appears.
    The report is saved in the same folder as your holdings file.
"""
import os
import sys
from datetime import date, timedelta

import pandas as pd

# ---- Settings ----
INPUT_FILE = None  # None = show a file-picker window; or paste a full path
COST_BASIS_IS_PER_SHARE = False
OPEN_REPORT_WHEN_DONE = True
LOOKBACK_DAYS = 10  # window used to find the latest close (covers weekends/holidays)


def read_holdings(path, per_share=False):
    df = pd.read_excel(path)
    df.columns = [str(c).strip().lower() for c in df.columns]

    required = ["ticker", "shares", "costbasis"]
    absent = [c for c in required if c not in df.columns]
    if absent:
        raise ValueError(
            f"Missing column(s): {', '.join(absent)}. "
            "Expected headers: Ticker, Shares, CostBasis"
        )

    df = df[required].copy()
    df["ticker"] = df["ticker"].astype(str).str.strip().str.upper()
    df = df[df["ticker"].notna() & (df["ticker"] != "") & (df["ticker"] != "NAN")]
    df["shares"] = pd.to_numeric(df["shares"], errors="coerce")
    df["costbasis"] = pd.to_numeric(df["costbasis"], errors="coerce")

    bad = df[df["shares"].isna() | df["costbasis"].isna()]
    if not bad.empty:
        raise ValueError(
            "Non-numeric Shares/CostBasis for: " + ", ".join(bad["ticker"])
        )

    if per_share:
        df["costbasis"] = df["costbasis"] * df["shares"]

    if df["ticker"].duplicated().any():
        print("Duplicate tickers found; combining lots.")
        df = df.groupby("ticker", as_index=False, sort=False).sum(numeric_only=True)
    return df.reset_index(drop=True)


def get_last_close(symbol):
    """Return (close_price, price_date) or (None, None) on failure."""
    import yfinance as yf

    try:
        hist = yf.Ticker(symbol).history(
            start=date.today() - timedelta(days=LOOKBACK_DAYS),
            auto_adjust=False,
        )
        close = hist["Close"].dropna()
        if close.empty:
            raise ValueError("no data returned")
        return float(close.iloc[-1]), close.index[-1].date()
    except Exception as exc:  # network errors, bad tickers, etc.
        print(f"WARNING: could not retrieve price for {symbol}: {exc}", file=sys.stderr)
        return None, None


def build_report(holdings, price_fn=get_last_close):
    quotes = [price_fn(t) for t in holdings["ticker"]]
    out = pd.DataFrame(
        {
            "Ticker": holdings["ticker"],
            "Shares": holdings["shares"],
            "CostBasis": holdings["costbasis"],
            "ClosePrice": [q[0] for q in quotes],
            "PriceDate": [q[1] for q in quotes],
        }
    )
    out["MarketValue"] = out["Shares"] * out["ClosePrice"]
    out["GainLoss"] = out["MarketValue"] - out["CostBasis"]
    out["GainLossPct"] = out["GainLoss"] / out["CostBasis"]
    out["PortfolioPct"] = out["MarketValue"] / out["MarketValue"].sum(skipna=True)
    return out


def portfolio_totals(report):
    # Only positions with a valid price count, so a failed lookup
    # doesn't silently distort the portfolio gain/loss.
    ok = report.dropna(subset=["MarketValue"])
    cost = ok["CostBasis"].sum()
    value = ok["MarketValue"].sum()
    gl = value - cost
    return {
        "cost": cost,
        "value": value,
        "gain_loss": gl,
        "gain_loss_pct": gl / cost if cost else float("nan"),
        "missing": report.loc[report["MarketValue"].isna(), "Ticker"].tolist(),
    }


def write_excel(report, totals, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"

    headers = list(report.columns)
    ws.append(headers)
    for row in report.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])

    n = len(report)
    last = n + 1
    tot = n + 3

    # Totals row uses live Excel formulas so the file stays auditable
    ws.cell(tot, 1, "PORTFOLIO TOTAL")
    # Cost total counts only positions that have a price, matching the console totals
    ws.cell(tot, 3, f'=SUMIFS(C2:C{last},F2:F{last},"<>")')
    ws.cell(tot, 6, f"=SUM(F2:F{last})")
    ws.cell(tot, 7, f"=SUM(G2:G{last})")
    ws.cell(tot, 8, f'=IF(C{tot}=0,"",G{tot}/C{tot})')
    ws.cell(tot + 2, 1, f"Prices as of most recent close; report generated {date.today()}")

    head_fill = PatternFill("solid", fgColor="1F3864")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for c in range(1, len(headers) + 1):
        ws.cell(tot, c).font = Font(bold=True)
        ws.cell(tot, c).border = Border(top=Side(style="thin"))

    for r in range(2, tot + 1):
        ws.cell(r, 2).number_format = "#,##0.####"
        for c in (3, 4, 6, 7):
            ws.cell(r, c).number_format = "$#,##0.00"
        for c in (8, 9):
            ws.cell(r, c).number_format = "0.00%"
        ws.cell(r, 5).number_format = "yyyy-mm-dd"

    green = PatternFill("solid", bgColor="C6EFCE")
    red = PatternFill("solid", bgColor="FFC7CE")
    rng = f"G2:H{tot}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                                  fill=green, font=Font(color="006100")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                  fill=red, font=Font(color="9C0006")))

    for col, w in zip("ABCDEFGHI", (18, 10, 14, 12, 12, 16, 14, 13, 13)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def choose_file():
    """Open a file-picker window and return the chosen path."""
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    path = filedialog.askopenfilename(
        title="Select your holdings spreadsheet",
        filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
    )
    root.destroy()
    return path


def main():
    # No terminal needed: run this file from your editor (Spyder, VS Code,
    # PyCharm, IDLE, Jupyter...) and pick the spreadsheet in the window.
    path = INPUT_FILE or choose_file()
    if not path:
        print("No file selected.")
        return
    output = os.path.join(os.path.dirname(os.path.abspath(path)), "portfolio_report.xlsx")

    holdings = read_holdings(path, per_share=COST_BASIS_IS_PER_SHARE)
    print(f"Fetching prices for {len(holdings)} tickers...")
    report = build_report(holdings)
    totals = portfolio_totals(report)

    pd.options.display.float_format = "{:,.2f}".format
    show = report.copy()
    show["GainLossPct"] = (show["GainLossPct"] * 100).map("{:.2f}%".format)
    show["PortfolioPct"] = (show["PortfolioPct"] * 100).map("{:.1f}%".format)
    print("\n" + show.to_string(index=False) + "\n")
    print(f"Total cost basis : ${totals['cost']:,.2f}")
    print(f"Portfolio value  : ${totals['value']:,.2f}")
    print(f"Total gain/loss  : ${totals['gain_loss']:,.2f} ({totals['gain_loss_pct']:.2%})")
    if totals["missing"]:
        print("\nWARNING: no price found for: " + ", ".join(totals["missing"])
              + " (excluded from totals)")

    try:
        write_excel(report, totals, output)
    except PermissionError:
        print("\nCould not save the report. If portfolio_report.xlsx is open in "
              "Excel, close it and run again.")
        return
    print(f"\nReport written to: {output}")
    if OPEN_REPORT_WHEN_DONE:
        try:
            os.startfile(output)  # Windows
        except AttributeError:
            import subprocess
            subprocess.run(["open" if sys.platform == "darwin" else "xdg-open", output])


if __name__ == "__main__":
    main()
